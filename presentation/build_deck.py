#!/usr/bin/env python3
"""Build the short, editable technical pitch with native Notes, video and placeholders."""
from __future__ import annotations
import json
import shutil
from lxml import etree
from PIL import Image
from pptx.util import Inches
from deck_style import *
from technical_slide import build_statistics, build_training
import deck_style as ds

ART = ROOT / 'assets/revision'
VIDEO = ROOT / 'assets/short/binder-reveal.mp4'
SOURCES['nsceb'] = ('NSCEB final report, 2025, §1.3', 'https://www.biotech.senate.gov/final-report/chapters/chapter-1/section-3/')
SOURCES['foresight'] = ('NSCEB final report, 2025, §3.3', 'https://www.biotech.senate.gov/final-report/chapters/chapter-3/section-3/')
SOURCES['aisi'] = ('UK AISI, Frontier AI Trends Report; models through October 2025', 'https://www.aisi.gov.uk/frontier-ai-trends-report')
SOURCES['classifiers'] = ('Anthropic, Constitutional Classifiers++', 'https://www.anthropic.com/research/next-generation-constitutional-classifiers')
SOURCES['davt'] = ('Pandeva et al., Deep anytime-valid hypothesis testing, AISTATS 2024', 'https://proceedings.mlr.press/v238/pandeva24a.html')
SOURCES['denoising'] = ('Vincent et al., Stacked Denoising Autoencoders, JMLR 2010', 'https://www.jmlr.org/papers/v11/vincent10a.html')


def pic(s, path, x, y, w, h):
    with Image.open(path) as im:
        ratio = im.width / im.height
    ww = min(w, h * ratio)
    return s.shapes.add_picture(str(path), Inches(x + (w - ww) / 2), Inches(y + (h - ww / ratio) / 2), width=Inches(ww), height=Inches(ww / ratio))


def heading(s, content, light=False, size=34):
    text(s, content, .62, 1.02, 12.1, 1.0, size, INK if light else PAPER, True)


def caption(s, content, x, y, w, light=False):
    text(s, content, x, y, w, .38, 11, DARKGREEN if light else MUTED)


def placeholder(s, x, y, w, h, name, detail, light=False):
    fill, stroke, fg, muted = ('E4E7DD', 'AABAAE', INK, DARKGREEN) if light else (PANEL, LINE, PAPER, GREEN)
    rect(s, x, y, w, h, fill, stroke)
    text(s, 'CAPTURE PLACEHOLDER', x+.18, y+.16, w-.36, .25, 10, muted, True, MONO)
    line(s, x+.17, y+.55, x+w-.17, y+.55, stroke)
    text(s, name, x+.2, y+.79, w-.4, .78, 23, fg, True)
    text(s, detail, x+.2, y+1.67, w-.4, max(.48,h-1.82), 15, muted)


def fade(s,shape):
    """Native PowerPoint entrance effect; exported PDFs show the final state."""
    xml=f'''<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="5" dur="800"/><p:tgtEl><p:spTgt spid="{shape.shape_id}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
    s._element.append(etree.fromstring(xml.encode()))


def architecture_cover(s):
    """Native editable architecture rendering; every visual element has a system role."""
    x0 = 6.58
    for x in [x0+i*.27 for i in range(24)]:
        for y in [1.18+i*.27 for i in range(19)]:
            dot(s,x,y,.009,LINE)
    text(s,'SOURCE → STATE → EXPERIMENT → REVIEW',6.55,1.03,6.02,.26,10,GREEN,True,MONO)
    for i in range(3):
        x,y=6.7+i*.20,1.73+i*.16
        rect(s,x,y,1.16,1.43,PANEL,LINE)
        for j in range(6): line(s,x+.15,y+.23+j*.16,x+.97,y+.23+j*.16,GREEN if j==0 else LINE,.8)
    text(s,'PAPERS',6.94,3.36,1.51,.28,11,MUTED,True,MONO)
    for a,b in [((8.05,2.59),(8.87,2.38)),((8.87,2.38),(9.72,1.95)),((8.87,2.38),(9.58,3.06)),((9.72,1.95),(10.30,2.53)),((9.58,3.06),(10.30,2.53)),((8.87,2.38),(8.59,3.10))]:
        line(s,*a,*b,GREEN,1.6)
    for x,y,r in [(8.87,2.38,.16),(9.72,1.95,.10),(9.58,3.06,.13),(10.30,2.53,.10),(8.59,3.10,.07)]:
        dot(s,x,y,r,GREEN)
    text(s,'CLAIM GRAPH',8.55,3.57,2.21,.28,11,GREEN,True,MONO)
    line(s,10.49,2.56,11.14,2.56,GREEN,1.4)
    line(s,11.14,2.56,11.14,4.18,GREEN,1.4)
    for y in [1.71,2.91,4.12]:
        line(s,11.14,y,11.62,y,GREEN,1.4)
        rect(s,11.62,y-.27,.82,.54,PANEL,GREEN,True)
        dot(s,11.83,y,.06,GREEN)
        line(s,11.97,y-.05,12.25,y-.05,MUTED,1)
        line(s,11.97,y+.07,12.20,y+.07,LINE,1)
    text(s,'AGENT SESSIONS',10.84,4.63,2.03,.50,11,GREEN,True,MONO)
    line(s,12.06,4.45,12.06,5.43,AMBER,1.4)
    line(s,12.06,5.43,10.67,5.43,AMBER,1.4)
    rect(s,8.53,5.06,2.12,.76,PANEL,AMBER)
    text(s,'CODE + DATA',8.73,5.29,1.77,.28,12,AMBER,True,MONO)
    line(s,8.51,5.43,7.88,5.43,AMBER,1.4)
    rect(s,6.66,5.06,1.2,.76,GREEN)
    text(s,'REVIEW',6.82,5.31,.95,.25,11,INK,True,MONO)
    line(s,7.23,5.03,7.23,4.25,GREEN,1.2,dash=True)
    line(s,7.23,4.25,9.41,4.25,GREEN,1.2,dash=True)
    line(s,9.41,4.25,9.41,3.37,GREEN,1.2,dash=True)
    text(s,'Persistent evidence + bounded compute',6.72,6.22,5.85,.42,17,MUTED)




def build():
    p = new_prs()
    s = slide(p, 1, 'Biological foresight')
    architecture_cover(s)
    text(s, 'Autonomous research\nfor biological\nforesight', .63, 1.28, 5.82, 2.47, 40, PAPER, True)
    text(s, 'Biological breakthroughs can create\nunilateral power before defenses adapt.', .67, 4.02, 5.79, 1.06, 23, GREEN)
    text(s, 'We connect the literature to executable\nexperiments and inspectable evidence.', .68, 5.53, 5.54, .98, 22, MUTED)
    note(s, 1, 'Autonomous research for biological foresight',
         'Biology is an underappreciated source of future unilateral power: a discovery can give one actor capabilities others have not anticipated. Reactive measures are inadequate when the capability itself is new. The national biotechnology commission describes the United States repeatedly playing catch-up. Our thesis is that scientists need systems that continuously connect published claims, formulate hypotheses and execute computational experiments. We are building that research engine, with a record of how every result was produced.', 30,
         'Strategic thesis and product ambition. Native editable rendering shows papers, claim graph, recursive sessions, executable experiments and review feedback. It is a conceptual system diagram, not a measured run. No operational threat prediction or biological prevention result is claimed.',
         'Working product name; edit freely. The rest of this deck puts technical mechanisms and measurements in the main argument. Preserve future-direction status where the complete capability has not been demonstrated.',
         'Problem & Real-World Impact (25); DTX strategic relevance', ['nsceb','foresight','dtx','ARCHITECTURE.md'])

    s = slide(p, 2, 'Experimental framing / why anticipatory research')
    heading(s, 'AI is accelerating biology. Defense needs a head start.', size=32)
    blocks = [
        ('CAPABILITY', 'Up to 60%', 'above PhD baselines on selected\nbiology / chemistry questions', 'AISI · open-ended QA benchmarks'),
        ('SAFEGUARDS', '198,000', 'adversarial attempts still exposed\none high-risk vulnerability', 'Anthropic · Constitutional Classifiers++'),
    ]
    for i,(label,metric,detail,source) in enumerate(blocks):
        x=.68+i*6.30
        line(s,x,2.22,x+5.61,2.22,LINE)
        text(s,label,x,2.42,5.56,.3,11,GREEN,True,MONO)
        text(s,metric,x,2.98,5.57,.83,45,PAPER,True)
        text(s,detail,x,4.00,5.57,.9,23,GREEN)
        caption(s,source,x,5.04,5.57)
    rect(s,.67,5.62,12.00,.99,PANEL)
    text(s,'“The United States tends to play catch-up after critical\ntechnologies have already become mainstream.”',.87,5.79,9.35,.65,20,PAPER)
    text(s,'NSCEB\n2025 · §1.3',10.83,5.88,1.55,.56,12,GREEN,True,MONO)
    note(s,2,'AI is accelerating biology. Defense needs a head start.',
         'Biology is becoming more accessible as AI learns to reason about it and execute longer tasks. AISI reports models scoring up to sixty percent above expert baselines on selected biology and chemistry questions. Safeguards are improving, but extensive adversarial testing still uncovers vulnerabilities. That does not show that a novice can build a weapon; it shows why safeguards alone are an incomplete preparedness strategy. Our proposed role is anticipatory research: test consequential scientific questions and assemble traceable evidence before an emerging capability becomes a crisis.',35,
         'AISI report covers models released through October2025. Relative QA score improvement, not a 60-percentage-point gain, whole-scientist competence or measured weapons feasibility. Private biology/chemistry sets each contain over280 open-ended questions; expert absolute scores roughly40–50%. Anthropic reports over1700 cumulative red-team hours and198000 attempts, one high-risk vulnerability, no universal jailbreak discovered in that evaluation. It is not a misuse probability. NSCEB §3.3 calls for emerging-biotechnology foresight and early detection/characterization. The proposed project contribution is an investment/product thesis, not validated threat prediction. No operational threat generation or jailbreak method is included.',
         'Experimental second slide requested by the user. Keep benchmark and red-team scope in the spoken explanation. The current system supports scientific investigation; the anticipatory-defense use case still needs domain-led evaluation.',
         'Problem & Real-World Impact (25); DTX relevance; Feasibility (25)', ['aisi','classifiers','nsceb','foresight'])

    s = slide(p, 3, 'Source-grounded recursive research')
    heading(s, 'A claim graph drives a persistent tree of experiments.')
    labels = [('Ingest → grounded graph', 'HGNC / ChEBI / MONDO → canonical IDs'), ('Recursive agent sessions', 'One action → execution → observation'), ('Acceptance + review', 'Typed checks / artifacts / human decision')]
    for i,(a,b) in enumerate(labels):
        x=.66+i*4.23
        text(s,a,x,2.18,3.86,.4,21,GREEN,True)
        text(s,b,x,2.71,3.89,.52,15,MUTED)
        if i<2: text(s,'→',x+3.84,2.29,.3,.4,20,GREEN)
    pic(s,ART/'knowledge.png',.65,3.30,5.83,3.02)
    pic(s,ART/'investigation.png',6.85,3.30,5.83,3.02)
    caption(s,'Knowledge: canonical identity, context and source inspection',.67,6.35,5.8)
    caption(s,'Investigation: branches, activity and experiment replay',6.87,6.35,5.8)
    text(s,'18 actions / checkpoint   ·   2,880 actions / shared tree   ·   atomic grants survive restart',.67,6.83,11.99,.21,11,GREEN,True,MONO)
    note(s, 3, 'A claim graph drives a persistent tree of experiments.',
         'The graph stores typed biological claims, resolved entity identifiers, context and source quotations. A resumable agent searches that graph and its papers, proposes hypotheses, writes analysis code and observes the result. It can fork a question into child sessions. Every eighteen research actions it reports to a parent controller, which continues, forks or prunes the branch. All descendants share a frozen budget; transactional grants and recorded decisions survive restart. The scientist sees the same provenance chain in the website, from a literature relationship to the branch and its experimental artifacts.', 40,
         'Implemented extraction, grounding, claim/evidence storage, recursive sessions, controller and runtime journal. The two images are actual current React UI captures at a054f1b with synthetic test data, not a measured six-agent investigation. Claim confidence counts distinct source papers, not independently deduplicated experiments. The controller has no measured equal-budget discovery advantage. Ordinary falsifier screens reported fields/flags, not an independent rerun of arbitrary code; human review records a decision.',
         'Replace both screenshots with the final corpus and selected investigation. Current sample data is marked on slide. For deeper questions: Opus extraction, separate Sonnet direction/repair sessions, closed-vocabulary grounding, DuckDB evidence and SQLite controller; the current executor uses host Python subprocesses, recorded submitted/executed code, scoped working paths and owned-process cancellation. This is not OS isolation; separate-account security remains deployment work.',
         'Technical Execution (50); AI Novelty (25); AI Sophistication (50); UX (40); Understanding User (30)', ['ARCHITECTURE.md','docs/claim-repair.md','docs/runtime.md','src/dnhacksbio/explorer/control.py','presentation/assets/revision/manifest.json'])
    tag(s,'WORKING WEBSITE · SAMPLE DATA',.71,3.36,3.31,AMBER,INK)

    s = slide(p, 4, 'Executable scientific instruments / grounding')
    heading(s, 'Images become evidence when checked against coordinates.',size=31)
    pic(s,REPO/'docs/inhibitor-review/agent-comparison.png',.65,2.15,7.35,4.37)
    tag(s,'RECORDED TOOL RUN',.69,2.18,2.25,GREEN,INK)
    text(s,'Docking → image review → coordinate countercheck',.7,6.48,7.26,.42,17,PAPER,True)
    movie = s.shapes.add_movie(str(VIDEO), Inches(8.36), Inches(2.15), Inches(4.3), Inches(2.42), poster_frame_image=str(REPO/'frontend/src/demo/captures/demo-expanded.png'), mime_type='video/mp4')
    for cond in s._element.xpath('./p:timing//p:video/p:cMediaNode/p:cTn/p:stCondLst/p:cond'):
        cond.set('delay','0')
    button=rect(s,8.37,4.69,2.71,.5,GREEN)
    button.click_action.hyperlink.address='binder-reveal.mp4'
    playlabel=text(s,'▶ OPEN VIDEO',8.56,4.78,2.31,.30,15,INK,True)
    playlabel.click_action.hyperlink.address='binder-reveal.mp4'
    text(s,'23 seconds',11.24,4.84,1.35,.24,11,MUTED)
    caption(s,'Auto-plays in Slide Show · MP4 in delivery folder',8.39,5.32,4.22)
    caption(s,'Binder cinematic: illustrative geometry',8.39,5.70,4.22)
    text(s,'MEASURED CONTROL',8.39,6.18,4.17,.26,11,AMBER,True,MONO)
    text(s,'2.675 Å RMSD  /  required ≤2 Å',8.39,6.55,4.25,.35,17,PAPER,True)
    note(s, 4, 'Images become evidence when checked against coordinates.',
         'These are scientific instruments the agent can operate, not just images returned by a chatbot. We ran three docking seeds. The agent inspected the paired molecular images and counterchecked the same contact using canonical coordinates. The selected contact was 3.132 versus 5.324 angstroms, but the whole-pose recovery error was 2.675 against a two-angstrom requirement. That control failed, and the failure stays in the record. The toolkit also includes PhysiCell and BioFVM tissue simulation, Cytosim spindle mechanics and binder-interface inspection. The short binder movie illustrates the website interaction; it is not a new molecular result.', 45,
         'Actual 3VQU/O22 workbench and recorded agent countercheck. Three-seed recovery criterion is RMSD ≤2 Å; proximity does not establish efficacy. The 23.44-second H.264 movie is the authored binder node-reveal browser capture with illustrative geometry. Tissue and spindle are conditional, uncalibrated simulations. Binder interface exists; no BindCraft inference pilot is claimed. Native movie start condition is delay=0; PowerPoint playback still needs local rehearsal.',
         'Play the deck in Slide Show mode: this movie is configured to start when the slide opens. The green button opens the adjacent binder-reveal.mp4 in the extracted delivery folder; the same clip was sent separately to Telegram. A PDF or Telegram document preview cannot play embedded PowerPoint media. Swap the cinematic for a final recorded real investigation when available.',
         'Technical Execution (50); AI Sophistication (50); Design Craft (30); Usability (40); AI Reliability (25)', ['docs/inhibitor-review/README.md','docs/inhibitor-review/agent-countercheck.json','docs/tumor-stroma.md','docs/spindle-simulator.md','docs/binder-design.md','docs/demo-cinematics.md'])

    build_training(p,5)
    build_statistics(p,6)

    s = slide(p, 7, 'Discovery demonstration / editable sequence')
    heading(s,'Trace an investigation, then reveal the later discovery.')
    text(s,'Source claim → untested bridge → executable analysis → result → checks → human review',.68,2.10,11.98,.67,23,GREEN,True)
    pic(s,ART/'investigation-detail.png',.67,2.99,6.09,3.42)
    tag(s,'WEBSITE PLACEHOLDER · SAMPLE RUN',.72,3.02,3.82,AMBER,INK)
    text(s,'[Proposed → passed → reviewed]  ·  [candidate + result]',.71,6.49,6.17,.41,16,GREEN,True)
    group=s.shapes.add_group_shape()
    rect(group,7.16,3.05,5.46,3.79,PAPER)
    text(group,'RESEARCH ARTICLE  /  LATER PUBLICATION',7.41,3.27,4.95,.28,10,DARKGREEN,True,MONO)
    line(group,7.41,3.73,12.37,3.73,'B8C6B9')
    text(group,'[Paper title matching\nthis recovered finding]',7.42,4.03,4.85,1.04,27,INK,True)
    text(group,'[Publication date · DOI]',7.44,5.31,4.76,.32,14,DARKGREEN)
    text(group,'MATCHED FINDING',7.44,5.99,4.74,.24,10,DARKGREEN,True,MONO)
    text(group,'[The specific result shared with the candidate]',7.44,6.31,4.77,.39,16,INK)
    fade(s,group)
    note(s,7,'Trace an investigation, then reveal the later discovery.',
         '[Fill the bracketed run details before presenting.] Start with the source relationship in the website. Show the agent turning a gap into a question, then open its saved analysis and raw output. Show what the automatic checks accepted and the scientist reviewed. Of [proposed] candidates, [passed] reached review. Now click to reveal the later paper: this candidate matches [the exact finding]. Explain the shared result, not just overlapping keywords. The audience has seen the working loop first; on the next slide, show how its accessible evidence was frozen.',35,
         'User confirms demonstration records are on a teammate machine and authorizes inferred structure with placeholders. Counts, candidate, paper identity and exact match are not locally verified. The white paper mockup is native editable text/shapes, not an invented publication. Current left image is actual website UI with synthetic test data. The paper group has an800ms on-click fade; it is visible in PDF final state.',
         'Replace the left panel with a video or snapshots from the same selected run, and fill the actual survivors and paper. Show the full loop, then survivors, then click the paper reveal. Advance directly to slide 8. The user will manually edit this working deck.',
         'Technical Originality/Execution (50); Problem & Impact (25); Reliability/Trustworthiness (25)', ['plans/PLAN-demo-slides.md','presentation/assets/revision/manifest.json'])

    s=slide(p,8,'Frozen evidence / immediately after reveal',light=True)
    heading(s,'Backtest discovery against a recorded evidence cutoff.',light=True,size=32)
    placeholder(s,.68,2.42,5.54,3.80,'Library + source reader','[Capture the actual accessible corpus]\n[Open the source claim and quotation]\n[Keep the excluded paper outside this view]',light=True)
    fields=[('FREEZE BEFORE EXECUTION','[Run ID · cutoff · corpus / data hashes]'),('EXCLUDE AND AUDIT','Target paper withheld; retain retrieval and tool logs'),('RECORD COMPUTATION','[GPU task · device · time] + saved code / outputs')]
    for i,(a,b) in enumerate(fields):
        y=2.53+i*1.30
        line(s,6.7,y,12.58,y,'ABBCAF')
        text(s,a,6.73,y+.19,5.81,.26,11,DARKGREEN,True,MONO)
        text(s,b,6.73,y+.61,5.80,.54,19,INK)
    text(s,'Next proof: sealed post-cutoff data and equal-budget comparisons against retrieval-only and unguided agents.',.7,6.56,11.9,.42,16,DARKGREEN)
    note(s,8,'Backtest discovery against a recorded evidence cutoff.',
         'Here is the backtest behind the reveal. Freeze the papers, data and access policy before starting, with the matching later paper held out. Run the ordinary research workflow and retain source references, submitted code, outputs and retrieval logs. Then compare the exact candidate finding with the held-out publication. GPU use belongs to a particular recorded model fit or scientific computation; fill that task and timing from the teammate run. Our locally documented L40S work trained the expression representation. The next stronger test uses sealed post-cutoff data and the same compute budget for a retrieval-only or unguided agent.',35,
         'Inferred demo structure authorized by the user; teammate records supply the placeholders. Curated local PDAC collection:100papers dated2007–2025 with inclusive2026-01-25curation boundary; this is not itself a run-access boundary. No actual Library/reader image or teammate recovery records are local. The independent L40S expression fit is supported by training.json; it must not be relabeled GPU-powered recovery. A frozen retrieval corpus cannot rule out pretrained model knowledge. A later matching paper does not by itself establish causality, novelty, or equal-budget discovery advantage. A real audit must check external retrieval and exact available data, including locally cached artifacts.',
         'Keep immediately after the paper-reveal slide. Fill every bracketed field from the selected run and exclusion audit. Replace the left native panel with an actual Library/source-reader screenshot. Do not turn a curation date into a claim about model memory or unrestricted retrieval.',
         'AI Reliability/Evaluation (25); Technical Originality (50); UX/Understanding User (40/30)', ['demo/pdac/README.md','demo/pdac/papers.json','plans/PLAN-demo-slides.md'])

    s=slide(p,9,'Deployment / commercial proof')
    heading(s,'Deploy first inside a computational research team.')
    text(s,'USER',.68,2.14,5.74,.27,11,GREEN,True,MONO)
    text(s,'A research lead deciding what to test\nwith limited compute and reviewer time.',.68,2.61,7.41,1.12,26,PAPER,True)
    steps=[('First integration','Approved corpus, data permissions, owned tool execution'),('Unit economics','All model + tool + review cost per rerunnable useful result'),('Adoption test','Review time, repeat use and paid expansion within the team')]
    for i,(a,b) in enumerate(steps):
        y=4.10+i*.77
        text(s,a,.70,y,3.04,.43,19,GREEN,True)
        text(s,b,3.75,y+.01,4.44,.56,17,MUTED)
    pic(s,REPO/'docs/tissue-review/native-final-Presentation-Comparison.png',8.61,2.12,4.00,2.33)
    pic(s,REPO/'docs/spindle-review/motors.png',8.61,4.77,4.00,1.72)
    caption(s,'Tissue: PhysiCell / BioFVM conditional simulation',8.64,4.48,3.99)
    caption(s,'Spindle: Cytosim mechanics / corrected motor export',8.64,6.58,3.99)
    note(s,9,'Deploy first inside a computational research team.',
         'The first customer is a computational research lead who already has questions, data and a limited experimental budget. Start in one private workspace and integrate with that team’s literature and analysis workflow. Measure independently rerunnable useful findings, total model and tool cost, and reviewer time. A subscription plus metered compute is a business model to test. The longer-term direction is an autonomous research organization: allocating experiments across scientific instruments, accumulating evidence and helping scientists anticipate consequential biological capabilities before they become widespread.',30,
         'Buyer and commercial model are hypotheses; no traction or pricing validation is claimed. Existing runtime records action/operation time and SDK usage; cost per useful discovery has not been measured. Research branch monitoring has prefix models/grouped calibration code but real-trajectory fitting remains pending. Separate-account deployment, execution permissions, data rights, privacy and security review remain deployment work. The tissue and spindle panels are actual conditional software outputs; biological calibration is not established.',
         'Add the team’s final ask and any measured pilot cost/user feedback. Future direction may be presented boldly; do not invent customers or effectiveness. Medical/clinical use would require additional validation and regulatory work. The active ecosystem/pharmacotype research routes were cancelled; these distinct 3D tools are not those routes.',
         'Feasibility & Deployment (25); Understanding User (30); Design Craft (30); Problem & Impact (25)', ['docs/deployment.md','docs/runtime.md','docs/branch-monitoring.md','docs/tumor-stroma.md','docs/spindle-motors.md','dtx'])

    s=slide(p,10,'Three-year direction / investment thesis',light=True)
    heading(s,'Stronger models should produce more science per research budget.',light=True,size=30)
    years=[
        ('YEAR 1 / 2027','A repeatable workflow','Connect one team’s sources and tools.\nEstablish equal-budget baselines.','Prove: rerunnable findings,\nreview time, cost and repeat use.'),
        ('YEAR 2 / 2028','Allocation that learns','Train on completed trajectories.\nCalibrate on independent episodes.','Prove: discovery yield under\na fixed budget; paid expansion.'),
        ('YEAR 3 / 2029','Continuous programs','Coordinate supervised portfolios\nas models and instruments improve.','Prove: retained customers,\nuseful output and scalable margins.'),
    ]
    for i,(year,title,body,proof) in enumerate(years):
        x=.69+i*4.22
        text(s,year,x,2.48,3.89,.28,11,DARKGREEN,True,MONO)
        line(s,x,3.00,x+3.81,3.00,'ACBCAE')
        text(s,title,x,3.25,3.81,.8,27,INK,True)
        text(s,body,x,4.42,3.80,1.02,18,DARKGREEN)
        text(s,proof,x,5.73,3.80,.85,18,INK,True)
    note(s,10,'Stronger models should produce more science per research budget.',
         'Our three-year thesis is that model progress expands this system’s useful workload. Better extraction fills the grounded graph; better coding and instrument use expand what an investigation can execute. We keep the experimental records and evaluation contracts while changing the model. First prove a repeatable team workflow. Then train allocation on completed trajectories and demonstrate better results at the same budget. By year three, the optimistic product is continuous, supervised research programs. The company must earn its position through trusted integrations, accumulated evidence and customer retention, as stronger models become widely available.',35,
         'Explicit prospective milestones, not forecasts or delivered capabilities. Real branch-monitor fitting, calibration and operational statistical stopping are pending. Prospective independent episodes and careful handling of censored/pruned branches are required; related children are not independent calibration units. Better models can improve power or productivity but do not guarantee discovery yield. Permissioned sources, instrument integrations and accepted research histories are potential durable assets only if customers value them. No pricing, margin, revenue, retention or pilot uplift is fabricated. Three-year horizons are relative to the2026hackathon.',
         'Replace goals with measured pilot metrics as they become available. A useful VC dashboard records all-in cost per independently rerunnable useful result, expert review minutes, return usage, paid expansion, retention and integration time. Compare models/workflows at equal budgets and retain negative outcomes.',
         'Feasibility & Deployment (25); Technical Originality (50); Problem & Impact (25)', ['docs/branch-monitoring.md','docs/deployment.md','docs/runtime.md','dtx'])

    invalid=[b for b in BOUNDS if b[1]<0 or b[2]<0 or b[1]+b[3]>W+.01 or b[2]+b[4]>H+.01]
    if invalid: raise ValueError(f'Text outside slide: {invalid}')
    path=OUT/'DNHacks_2026_DN_Research.pptx'
    p.save(path)
    shutil.copy2(VIDEO,OUT/'binder-reveal.mp4')
    seconds=sum(n['seconds'] for n in TALK)
    (ROOT/'speaker-notes.md').write_text('# DNHacks technical pitch\n\n'+str(len(p.slides))+' slides, '+str(seconds)+' scripted seconds. Every slide has editable PowerPoint Notes.\n\n'+'\n\n'.join(f'## {n["slide"]:02} — {n["title"]}\n\n{n["notes"]}' for n in TALK)+'\n')
    manifest={'source_snapshot':ds.SNAPSHOT,'website_capture_snapshot':'a054f1b133a5c42583da4f9efd6311a70782b0da','main_slides':len(p.slides),'core_slides':list(range(1,len(p.slides)+1)),'core_seconds':seconds,'hidden_slides':[],'paper_reveal_slot':7,'frozen_corpus_slot':8,'video_slide':4,'video_start':'automatic; delay=0','external_video':'binder-reveal.mp4','fonts':[FONT,MONO]}
    (OUT/'deck-manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print(json.dumps(manifest,indent=2))


if __name__=='__main__': build()
