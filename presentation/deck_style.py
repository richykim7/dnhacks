#!/usr/bin/env python3
"""Build an editable DNHacks pitch and matching layout kit. Requires python-pptx."""
from __future__ import annotations

import json
import math
import random
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
W, H = 13.333333, 7.5
INK, PAPER, GREEN, AMBER, ROSE = "17201F", "F1EFE7", "A6C9B8", "DEB577", "D99E99"
MUTED, LINE, PANEL, DARKGREEN = "AFB9B2", "3A4944", "222E2A", "365E4D"
FONT, MONO = "IBM Plex Sans", "IBM Plex Mono"
SNAPSHOT = "fee316580ba27951c2ad6fea4bdd4a6b335421ed"
SOURCES = {
    "pubmed": ("NIH / NLM · About PubMed", "https://pubmed.ncbi.nlm.nih.gov/about/"),
    "seer": ("NCI SEER · Pancreatic Cancer Stat Facts", "https://seer.cancer.gov/statfacts/html/pancreas.html"),
    "dtx": ("DTX Ventures · Vision", "https://www.dtxventures.com/vision"),
    "event": ("DNHacks · Official event and category descriptions", "https://dnhacks.org/"),
    "elicit": ("Elicit · Introducing Elicit Research Agent", "https://elicit.com/blog/introducing-elicit-research-agent"),
    "deepmind": ("Google DeepMind · Co-scientist", "https://deepmind.google/blog/co-scientist-a-multi-agent-ai-partner-to-accelerate-research/"),
}
TALK = []
BOUNDS = []


def color(v):
    return RGBColor.from_string(v)


def flat(sh):
    """Suppress the template's inherited Office drop-shadow preset."""
    for ref in sh._element.xpath('.//a:effectRef'):
        ref.set('idx', '0')
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh


def rect(s, x, y, w, h, fill, stroke=None, radius=False):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color(fill)
    if stroke:
        sh.line.color.rgb = color(stroke)
        sh.line.width = Pt(.7)
    else:
        sh.line.fill.background()
    if radius:
        sh.adjustments[0] = .08
    return flat(sh)


def text(s, txt, x, y, w, h, size=22, c=PAPER, bold=False, font=FONT,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, link=None):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = valign
    for i, line in enumerate(txt.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = 1.08
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color(c)
        if link:
            r.hyperlink.address = link
    BOUNDS.append((txt[:60], x, y, w, h))
    return sh


def line(s, x1, y1, x2, y2, c=LINE, width=1, arrow=False, dash=False):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = color(c)
    sh.line.width = Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        e = OxmlElement("a:tailEnd")
        e.set("type", "triangle")
        sh.line._get_or_add_ln().append(e)
    return flat(sh)


def dot(s, x, y, r=.045, c=GREEN, stroke=None):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-r), Inches(y-r), Inches(r*2), Inches(r*2))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color(c)
    if stroke:
        sh.line.color.rgb = color(stroke)
    else:
        sh.line.fill.background()
    return flat(sh)


def tag(s, txt, x, y, w=2.2, c=GREEN, bg=PANEL):
    rect(s, x, y, w, .31, bg)
    text(s, txt, x+.10, y+.052, w-.18, .21, 9.4, c, True, MONO)


def new_prs():
    p = Presentation()
    p.slide_width, p.slide_height = Inches(W), Inches(H)
    p.core_properties.title = "DN Research — DNHacks 2026"
    p.core_properties.subject = "Editable working pitch: scientific discovery, AI engineering and trustworthy evidence"
    p.core_properties.author = "DNHacks team"
    p.core_properties.keywords = "DNHacks, DN Research, working draft, scientific discovery"
    # Theme fonts keep newly inserted PowerPoint objects consistent with the deck.
    for rel in p.slide_master.part.rels.values():
        if rel.reltype.endswith('/theme'):
            from lxml import etree
            root = etree.fromstring(rel.target_part.blob)
            ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for node in root.xpath('//a:majorFont/a:latin | //a:minorFont/a:latin', namespaces=ns):
                node.set('typeface', FONT)
            rel.target_part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    return p


def slide(p, num, section, title=None, light=False, sub=None):
    s = p.slides.add_slide(p.slide_layouts[6])
    bg, fg, secondary = (PAPER, INK, DARKGREEN) if light else (INK, PAPER, MUTED)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(bg)
    text(s, "DN / RESEARCH", .55, .27, 2.6, .24, 11, fg, True, MONO)
    text(s, section.upper(), 5.0, .28, 7.76, .22, 10, secondary, False, MONO, PP_ALIGN.RIGHT)
    line(s, .55, .72, 12.78, .72, "CDCFC4" if light else LINE, .7)
    line(s, .55, 7.04, 12.78, 7.04, "CDCFC4" if light else LINE, .7)
    text(s, "DNHACKS 2026  /  WORKING PRESENTATION", .55, 7.16, 8, .19, 8, secondary, False, MONO)
    text(s, f"{num:02}", 12.14, 7.11, .64, .29, 12, fg, False, MONO, PP_ALIGN.RIGHT)
    if title:
        text(s, title, .55, 1.02, 12.0, 1.32, 36, fg, True)
    if sub:
        text(s, sub, .57, 2.13, 11.8, .55, 18, secondary)
    return s


def note(s, num, title, script, secs, proof="", update="", criteria="", sources=()):
    source_lines = []
    for source in sources:
        if source in SOURCES:
            name, url = SOURCES[source]
            source_lines.append(f"{name}: {url} (accessed 2026-09-06)")
        else:
            source_lines.append(f"Repository: {source} at {SNAPSHOT}")
    body = (f"TENTATIVE SCRIPT · {secs} seconds\n{script}\n\n"
            f"EVIDENCE / STATUS\n{proof}\n\n"
            f"BEFORE PRESENTING / NEXT EDIT\n{update or 'Recheck against the final demo build; preserve evidence labels.'}\n\n"
            f"JUDGING CRITERIA\n{criteria}\n\nSOURCES\n" + "\n".join(source_lines))
    s.notes_slide.notes_text_frame.text = body
    TALK.append({"slide": num, "title": title, "seconds": secs, "notes": body})


def source_footer(s, txt, key=None, light=False):
    # Full source URLs remain in speaker notes; footers keep the deck palette.
    text(s, txt, .57, 6.67, 12.1, .25, 9, DARKGREEN if light else MUTED)


def graph(s, x, y, w, h, seed=42, pale=False):
    rng = random.Random(seed)
    clusters = [(.20,.30),(.74,.21),(.60,.77)]
    nodes = []
    for k,(cx,cy) in enumerate(clusters):
        group=[]
        for j in range(20):
            a=rng.random()*math.tau
            rr=rng.uniform(.08,.23)
            xx=x+(cx+math.cos(a)*rr)*w
            yy=y+(cy+math.sin(a)*rr)*h
            group.append((xx,yy))
        nodes.append(group)
    for group in nodes:
        for j,a in enumerate(group):
            closest=sorted(range(len(group)),key=lambda k: math.dist(a,group[k]))[1:4]
            for k in closest:
                if k>j:
                    b=group[k]
                    line(s,*a,*b,"BFC9BE" if pale else "394D44",.65)
    for i,group in enumerate(nodes):
        for j,(xx,yy) in enumerate(group):
            dot(s,xx,yy,.025 if j%4 else .045,"698E7B" if pale else "6D9881")
    selected = [nodes[0][2],nodes[0][9],nodes[1][3],nodes[2][2],nodes[2][13]]
    for a,b in zip(selected,selected[1:]):
        line(s,*a,*b,DARKGREEN if pale else GREEN,2,dash=(a==selected[1]))
    for xx,yy in selected:
        dot(s,xx,yy,.095,PAPER if pale else INK,DARKGREEN if pale else GREEN)
        dot(s,xx,yy,.04,DARKGREEN if pale else GREEN)


def card(s,x,y,w,num,title,body,c=GREEN):
    line(s,x,y,x+w,y,c,1.6)
    text(s,num,x,y+.18,w,.4,16,c,False,MONO)
    text(s,title,x,y+.72,w,1.0,27,PAPER,True)
    text(s,body,x,y+1.88,w,1.2,19,MUTED)


def templates():
    p=new_prs()
    examples=[("Evidence slide","[One finding in a complete sentence]","[Claim and scope]","[Editable visual: chart, graph or source excerpt]"),
              ("Result slide","[What changed after the experiment?]","[Method · units · comparator]","[Measured result and uncertainty]"),
              ("Demo slide","[One question. One decision.]","[Actual run · date · dataset]","[Replace this panel with the real demo capture]"),
              ("Roadmap slide","[The next milestone earns the next claim]","[Implemented]","[Next / proposed] / [Later / planning]")]
    for i,(kind,title,small,body) in enumerate(examples,1):
        s=slide(p,i,"Layout kit / "+kind,title,light=(i==2))
        if i==1:
            tag(s,small,.59,2.47,5.2)
            rect(s,.6,3.13,8.03,3.11,PANEL,LINE)
            text(s,body,1.01,4.13,7.19,1.16,29,GREEN,True)
            text(s,"[Why it matters]",9.16,3.16,3.5,.64,27,PAPER,True)
            text(s,"[One short explanation.\nPut detail in speaker notes.]",9.17,4.12,3.52,1.56,22,MUTED)
        elif i==2:
            tag(s,small,.59,2.47,5.2,DARKGREEN,"DFE5DB")
            line(s,6.46,3.18,6.46,5.96,"B5C0B5")
            for x,label in [(.60,"[Baseline]"),(7.02,"[Our method]")]:
                text(s,label,x,3.30,5.42,.52,24,DARKGREEN,True)
                text(s,"[Value]",x,4.16,5.44,1.10,62,INK,True)
                text(s,"[Units · sample size · uncertainty]",x,5.53,5.44,.63,19,DARKGREEN)
        elif i==3:
            tag(s,small,.59,2.20,5.2)
            rect(s,.6,2.93,12.12,3.29,PANEL,LINE)
            text(s,body,1.06,4.02,11.23,1.08,32,GREEN,True,align=PP_ALIGN.CENTER)
        else:
            for j,(a,b) in enumerate([("NOW / IMPLEMENTED","[What works today]"),("NEXT / PROPOSED","[The proof milestone]"),("LATER / PLANNING","[The expansion]")]):
                card(s,.59+j*4.17,2.70,3.78,a,b,"[One concrete deliverable]\n[Its evidence requirement]\n[Owner or dependency]",GREEN if j==0 else AMBER)
        source_footer(s,"[Source, measured/planned/illustrative status, scope and date]",light=(i==2))
        s.notes_slide.notes_text_frame.text=("TENTATIVE SCRIPT\n[Write 40–65 words: finding, evidence, implication, transition.]\n\n"
          "STATUS\n[Implemented / measured / planned / illustration]\n\nSOURCE\n[Exact URL or repository path, version and date]\n\n"
          "EDITING\nDuplicate this slide; replace bracketed text and the visual panel. Preserve 0.55-inch margins, 36pt headline, 19–24pt body, IBM Plex fonts and explicit status. Avoid exceeding one argument per slide.")
    p.save(OUT/'DNHacks_Layout_Kit.pptx')
