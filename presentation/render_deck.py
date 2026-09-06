#!/usr/bin/env python3
"""Render and inspect the delivery packages using local LibreOffice and Poppler."""
from __future__ import annotations
import hashlib
from io import BytesIO
import json
import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from pathlib import Path
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'output'


def run(*args):
    return subprocess.run(args,check=True,capture_output=True,text=True).stdout


def words(s):
    return set(re.findall(r'\w+',unicodedata.normalize('NFKC',s).casefold()))


def shape_text(shapes):
    result=[]
    for s in shapes:
        if s.has_text_frame: result.append(s.text)
        if s.shape_type==6: result.extend(shape_text(s.shapes))
    return result


def render():
    report={'renderer':'LibreOffice Impress + Poppler','native_powerpoint_animation_review':False,'native_powerpoint_media_playback_review':False,'decks':{}}
    with tempfile.TemporaryDirectory(prefix='dnhacks-slide-review-') as tmp:
        tmp=Path(tmp)
        for name in ['DNHacks_2026_DN_Research']:
            source=OUT/(name+'.pptx')
            with zipfile.ZipFile(source) as z:
                assert z.testzip() is None
            prs=Presentation(source)
            assert all(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip() for s in prs.slides)
            full=tmp/(name+'-all'); full.mkdir()
            visible=Presentation(source)
            for item in visible.slides:
                item._element.set('show','1')
                # Impress may render an embedded movie as a gray box. Use its exact
                # poster in the PDF-only copy; retain the movie in the delivered PPTX.
                for shape in list(item.shapes):
                    if not shape._element.xpath('.//a:videoFile'):
                        continue
                    rel = shape._element.xpath('.//a:blip')[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    poster = shape.part.related_part(rel).blob
                    original = shape._element
                    parent = original.getparent()
                    index = list(parent).index(original)
                    still = item.shapes.add_picture(BytesIO(poster), shape.left, shape.top, shape.width, shape.height)
                    parent.remove(original)
                    parent.remove(still._element)
                    parent.insert(index, still._element)
                for timing in item._element.xpath('./p:timing'):
                    item._element.remove(timing)
            for item in visible._element.xpath('./p:custShowLst'): visible._element.remove(item)
            review_source=full/source.name
            visible.save(review_source)
            run('libreoffice',f'-env:UserInstallation={(tmp/(name+"-full-profile")).as_uri()}','--headless','--convert-to','pdf','--outdir',str(full),str(review_source))
            allpdf=full/(name+'.pdf')
            info=run('pdfinfo',str(allpdf))
            page_count=int(re.search(r'Pages:\s+(\d+)',info).group(1))
            if page_count!=len(prs.slides): raise ValueError(f'{name}: {page_count} rendered vs {len(prs.slides)} slides')
            pdf_text=run('pdftotext','-layout',str(allpdf),'-').split('\f')
            missing={}
            for i,s in enumerate(prs.slides):
                native=' '.join(shape_text(s.shapes))
                # Line-breaking can join/split number labels in either renderer; compare textual words.
                absent=words(native)-words(pdf_text[i])
                absent={w for w in absent if not w.isdigit() and len(w)>1}
                if absent: missing[str(i+1)]=sorted(absent)
            shutil.copy2(allpdf,OUT/(name+'.pdf'))
            review_name='DNHacks_2026_Reviewer_Copy.pdf' if name=='DNHacks_2026_DN_Research' else name+'.pdf'
            shutil.copy2(allpdf,OUT/review_name)
            run('pdftoppm','-scale-to','960','-png',str(allpdf),str(full/'slide'))
            pages=sorted(full.glob('slide-*.png'))
            thumbs=[]
            for i,p in enumerate(pages):
                im=Image.open(p).convert('RGB'); im.thumbnail((480,270))
                tile=Image.new('RGB',(500,305),'#D9DDD3');tile.paste(im,((500-im.width)//2,10))
                ImageDraw.Draw(tile).text((12,283),f'{i+1:02}',fill='#17201F')
                thumbs.append(tile)
            cols=2
            sheet=Image.new('RGB',(cols*500,((len(thumbs)+cols-1)//cols)*305),'#D9DDD3')
            for i,im in enumerate(thumbs):sheet.paste(im,((i%cols)*500,(i//cols)*305))
            sheet.save(OUT/('contact-sheet.png' if name=='DNHacks_2026_DN_Research' else name+'-contact-sheet.png'))
            preview=OUT/'review-pages'/name
            preview.mkdir(parents=True,exist_ok=True)
            for old in preview.glob('slide-*.png'):old.unlink()
            for f in pages:shutil.copy2(f,preview/f.name)
            if missing: raise ValueError(f'{name}: missing rendered text: {missing}')
            with zipfile.ZipFile(source) as package:
                movies=[n for n in package.namelist() if n.startswith('ppt/media/') and n.endswith('.mp4')]
            report['decks'][name]={'embedded_movies':len(movies),'slides':len(prs.slides),'notes':len(prs.slides),'rendered_pages':page_count,'hidden_slides':[i+1 for i,s in enumerate(prs.slides) if s._element.get('show')=='0'],'missing_text_tokens':missing}
        report['sha256']={f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in OUT.glob('*') if f.suffix in ('.pptx','.pdf')}
        (OUT/'artifact-checks.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))


if __name__=='__main__':render()
