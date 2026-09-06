"""Public Fudan external benchmark acquisition and strict sample/annotation joins."""
from __future__ import annotations

from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import re
import urllib.request

import numpy as np
import pandas as pd

from .protein_data import read_matrix, write_json
from .protein_validation import MODULES

BASE = "https://media.springernature.com/original/springer-static/esm/art%3A10.1186%2Fs13045-022-01384-3/MediaObjects/"
SOURCES = {f"supplement-{n}.xlsx": BASE + f"13045_2022_1384_MOESM{n}_ESM.xlsx" for n in (23, 25)}
SOURCES["hgnc.tsv"] = "https://storage.googleapis.com/public-download-files/hgnc/tsv/tsv/hgnc_complete_set.txt"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def acquire(raw):
    raw = Path(raw)
    raw.mkdir(parents=True, exist_ok=True)
    receipts = []
    for name, url in SOURCES.items():
        path = raw / name
        if not path.exists():
            with urllib.request.urlopen(url, timeout=120) as source:
                data = source.read(100 * 1024**2 + 1)
            if len(data) > 100 * 1024**2:
                raise ValueError("Source exceeds 100 MiB budget")
            if not data.startswith(b"PK" if name.endswith("xlsx") else b"hgnc_id"):
                raise ValueError("Unexpected source format")
            part = path.with_suffix(path.suffix + ".part")
            part.write_bytes(data)
            part.rename(path)
        receipts.append({"file": name, "url": url, "bytes": path.stat().st_size, "sha256": sha(path)})
    return receipts


def unique(values, kind):
    if any(not isinstance(x, str) or not x for x in values) or len(values) != len(set(values)):
        raise ValueError(f"Invalid or duplicated {kind}")
    return values


def read_fudan(raw):
    import openpyxl
    raw = Path(raw)
    book = openpyxl.load_workbook(raw / "supplement-23.xlsx", read_only=True, data_only=True)
    clinical = list(book["Table S1B"].values)
    header = clinical[0]
    unique([r[0] for r in clinical[1:]], "patients")
    metadata = {r[0]: dict(zip(header, r)) for r in clinical[1:]}
    runs = [(r[9], r[10]) for r in list(book["Table S1A "].values)[2:] if r[9]]
    unique([r[0] for r in runs], "samples")
    unique([r[1] for r in runs], "experiments")
    runs = dict(runs)
    book.close()
    book = openpyxl.load_workbook(raw / "supplement-25.xlsx", read_only=True, data_only=True)
    rows = book["Table S3A "].iter_rows(values_only=True)
    header = next(rows)
    columns = [(i, v) for i,v in enumerate(header) if isinstance(v, str) and re.fullmatch(r"PDAC_\d+T", v)]
    samples = unique([v for _,v in columns], "measured tumors")
    donors = unique([v[:-1] for v in samples], "measured tumor patients")
    if not set(samples) <= runs.keys() or not set(donors) <= metadata.keys():
        raise ValueError("Unjoined measured sample")
    symbols, data = [], []
    for row in rows:
        if not row[0]:
            continue
        symbols.append(row[0])
        values = []
        for i,_ in columns:
            v = row[i]
            if v in (None, "NA", "", "NaN"):
                values.append(float("nan"))
            elif isinstance(v, (int, float)) and np.isfinite(v) and v >= 0:
                values.append(float(v) if v > 0 else float("nan"))
            else:
                raise ValueError("Invalid FOT abundance")
        data.append(values)
    book.close()
    unique(symbols, "protein symbols")
    matrix = pd.DataFrame(np.asarray(data, dtype=np.float32).T, index=donors, columns=symbols)
    labels = []
    for donor, sample in zip(donors, samples):
        m = metadata[donor]
        grade = {"Poorly": "G3", "Moderately": "G2", "Well": "G1", "NA": None, None: None}.get(m["Differentiation"], "INVALID")
        if grade == "INVALID":
            raise ValueError("Unsupported differentiation label")
        purity = m["Tumor purity (%)"]
        labels.append({"donor": donor, "sample": sample, "experiment": runs[sample], "grade": grade,
                       "purity": float(purity) if isinstance(purity, (float, int)) else None})
    counts = Counter(r["grade"] or "unknown" for r in labels)
    audit = {"clinical_patients": len(metadata), "measured_unique_tumors": len(donors),
             "protein_symbols": len(symbols), "grades": dict(counts),
             "grade_pairs": min(counts["G3"], counts["G1"] + counts["G2"]),
             "sample_run_join": "one-to-one, complete", "treatment": "treatment-naive per primary publication",
             "canonical_identity": "published patient/sample/experiment crosswalk; no repeated tumors",
             "processing": "published FOT abundances; cohort-wide 1/6 detection filtering",
             "status": "external benchmark, previously analyzed publication cohort; not fresh confirmation"}
    return matrix, labels, audit


def annotation(raw):
    table = pd.read_csv(Path(raw) / "hgnc.tsv", sep="\t", dtype=str).fillna("")
    by_ensembl, by_name = defaultdict(set), defaultdict(set)
    for _, r in table.iterrows():
        if r["status"] != "Approved":
            continue
        if re.fullmatch(r"ENSG\d+", r["ensembl_gene_id"]):
            by_ensembl[r["ensembl_gene_id"]].add(r["symbol"])
        by_name[r["symbol"]].add(r["symbol"])
        for name in r["prev_symbol"].split("|"):
            if name:
                by_name[name].add(r["symbol"])
    return ({k: next(iter(v)) for k,v in by_ensembl.items() if len(v)==1},
            {k: next(iter(v)) for k,v in by_name.items() if len(v)==1})


def prepare(raw, cptac, curated, output):
    raw, cptac, curated, output = map(Path, (raw, cptac, curated, output))
    output.mkdir(parents=True, exist_ok=False)
    receipts = acquire(raw)
    external, labels, audit = read_fudan(raw)
    gene_map, symbol_map = annotation(raw)
    sources = {c: read_matrix(cptac / f"{c}_proteomics.txt") for c in ("BRCA", "COAD", "LUAD", "PDAC")}
    feature_ids = sorted(set().union(*(set(t.columns) for t in sources.values())))
    by_symbol = defaultdict(list)
    for feature in feature_ids:
        symbol = gene_map.get(feature.split(".")[0])
        if symbol:
            by_symbol[symbol].append(feature)
    external_by_symbol = defaultdict(list)
    for symbol in external.columns:
        if symbol in symbol_map:
            external_by_symbol[symbol_map[symbol]].append(symbol)
    shared = sorted(s for s,ids in by_symbol.items() if len(ids)==1 and len(external_by_symbol[s])==1)
    if len(shared) < 100:
        raise ValueError("Insufficient unambiguous cross-assay feature bridge")
    ids = [by_symbol[s][0] for s in shared]
    matrices = {c: t.reindex(columns=ids).to_numpy(dtype=np.float32) for c,t in sources.items()}
    dev = json.loads((curated / "dev.json").read_text())
    grades = {r["donor_id"]: r["grade"] for r in dev["observations"]}
    y = np.asarray([1 if grades[d]=="G3" else 0 if grades[d] in {"G1", "G2"} else -1 for d in sources["PDAC"].index])
    # External outcomes are deliberately a separate artifact, excluded from the GPU training bundle.
    np.savez_compressed(output / "inputs.npz", train=np.concatenate((matrices["BRCA"], matrices["COAD"])),
                        validation=matrices["LUAD"], development=matrices["PDAC"], development_y=y,
                        external=external[[external_by_symbol[s][0] for s in shared]].to_numpy(dtype=np.float32),
                        symbols=np.asarray(shared,dtype=str), development_ids=np.asarray(sources["PDAC"].index,dtype=str),
                        external_ids=np.asarray(external.index,dtype=str))
    write_json(output / "external-labels.json", labels)
    audit.update(sources=receipts, common_unambiguous_genes=len(shared),
                 ambiguous_cptac_symbols=sum(len(v)>1 for v in by_symbol.values()),
                 modules=MODULES, inputs_sha256=sha(output / "inputs.npz"),
                 source_matrices={c: sha(cptac / f"{c}_proteomics.txt") for c in sources})
    write_json(output / "audit.json", audit)
    return audit


def summarize(root, predictions, output):
    """Evaluate already-frozen predictions; never fit or choose a sign/threshold."""
    from scipy.stats import spearmanr
    from sklearn.metrics import roc_auc_score, balanced_accuracy_score
    root, predictions = Path(root), Path(predictions)
    training = json.loads((predictions.parent / "training.json").read_text())
    if (training["predictions_sha256"] != sha(predictions) or
        training["model_sha256"] != sha(predictions.parent / "model.npz") or
        training["inputs_sha256"] != sha(root / "inputs.npz") or
        training["external_labels_accessed"] is not False):
        raise ValueError("Frozen prediction/model/input provenance mismatch")
    labels = json.loads((root / "external-labels.json").read_text())
    lookup = {r["donor"]: r for r in labels}
    with np.load(predictions, allow_pickle=False) as archive:
        donors = unique(archive["donors"].tolist(), "prediction donors")
        if set(donors) != set(lookup):
            raise ValueError("Prediction/label donor mismatch")
        rows = [lookup[d] for d in donors]
        y = np.asarray([1 if r["grade"]=="G3" else 0 if r["grade"] in {"G1","G2"} else -1 for r in rows])
        coverage = archive["coverage"]
        purity = np.asarray([r["purity"] if r["purity"] is not None else np.nan for r in rows])
        if coverage.shape != y.shape or not np.isfinite(coverage).all():
            raise ValueError("Invalid coverage")
        masks = {"all_graded": y>=0, "coverage_at_least_80pct": (y>=0)&(coverage>=.8),
                 "purity_below_50pct": (y>=0)&(purity<50), "purity_at_least_50pct": (y>=0)&(purity>=50)}
        results = {}
        for name in ("rank_pca","rank_linear","fixed_modules","missingness","rank_kernel"):
            score = archive[name]
            if score.shape != y.shape or not np.isfinite(score).all():
                raise ValueError("Invalid frozen predictions")
            groups = {}
            for group,mask in masks.items():
                a,b = y[mask],score[mask]
                counts = np.bincount(a,minlength=2)
                record = {"n":len(a),"G1_G2":int(counts[0]),"G3":int(counts[1]),"pairs":int(counts.min())}
                if counts.min()>0:
                    rng=np.random.default_rng(20260906)
                    zero,one=np.flatnonzero(a==0),np.flatnonzero(a==1)
                    boots=[]
                    for _ in range(2000):
                        ii=np.concatenate((rng.choice(zero,len(zero),replace=True),rng.choice(one,len(one),replace=True)))
                        boots.append(roc_auc_score(a[ii],b[ii]))
                    record.update(auroc=float(roc_auc_score(a,b)),bootstrap_ci95=np.quantile(boots,[.025,.975]).tolist(),
                                  balanced_accuracy=float(balanced_accuracy_score(a,b>0)))
                groups[group]=record
            def correlation(x):
                mask=np.isfinite(x)&(y>=0)
                if np.unique(x[mask]).size<2 or np.unique(score[mask]).size<2:
                    return None
                return float(spearmanr(x[mask],score[mask]).statistic)
            results[name]={"groups":groups,"score_coverage_spearman":correlation(coverage),
                           "score_purity_spearman":correlation(purity)}
    report={"schema":"ProteinExternalBenchmark-v1","cohort":"Fudan Tong2022",
            "primary_view":"rank_pca","results":results,"predictions_sha256":sha(predictions),
            "labels_sha256":sha(root/"external-labels.json"),"bootstrap_replicates":2000,
            "training_record_sha256":sha(predictions.parent/"training.json"),
            "coverage_quantiles":np.quantile(coverage,[0,.25,.5,.75,1]).tolist(),
            "interpretation":"external transfer benchmark; cohort-wide source preprocessing and prior publication preclude untouched native confirmation",
            "incompatible_views":["original reference-abundance PCA", "original reference-abundance denoiser"]}
    write_json(output,report)
    return report
